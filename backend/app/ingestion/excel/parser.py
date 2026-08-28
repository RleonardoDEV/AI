"""Parseo de Excel guiado por contrato de datos.

Decisión ADR-004: no se usa la Excel REST API de Microsoft Graph, porque sus
páginas de referencia declaran `Files.ReadWrite` como permiso de mínimo
privilegio para operaciones de lectura. Se descarga el fichero y se parsea aquí
en modo solo lectura.

Principio (docs/00, P4): las filas que no validan no se descartan en silencio.
Van a cuarentena con su motivo y aparecen en el informe de calidad.
"""

from __future__ import annotations

import io
from collections.abc import Iterator, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ValidationError

from app.ingestion.excel.coercion import COERCERS, CoercionError, is_blank, normalize_key
from app.ingestion.excel.mapping import SchemaMapping, SchemaMappingError
from app.schemas import domain

# Longitud máxima del valor crudo que se guarda en cuarentena. Se conserva para
# que el PM pueda corregir el fichero; nunca se escribe en los logs (docs/04 §4).
_MAX_RAW_VALUE = 120

_MODELS: dict[str, type[BaseModel]] = {
    "TaskRow": domain.TaskRow,
    "RiskRow": domain.RiskRow,
    "ResourceRow": domain.ResourceRow,
    "BudgetRow": domain.BudgetRow,
}


@dataclass(frozen=True, slots=True)
class QuarantinedRow:
    """Una fila que no pudo convertirse al modelo canónico."""

    row_number: int
    reason: str
    field_name: str | None
    raw_value: str | None
    detail: str | None = None


@dataclass(frozen=True, slots=True)
class ParseResult:
    mapping_name: str
    sheet_name: str
    header_row: int
    rows: tuple[BaseModel, ...]
    quarantined: tuple[QuarantinedRow, ...]
    column_map: Mapping[str, str]
    missing_columns: tuple[str, ...]
    unmapped_headers: tuple[str, ...]

    @property
    def total_rows(self) -> int:
        return len(self.rows) + len(self.quarantined)

    @property
    def quarantine_ratio(self) -> float:
        return len(self.quarantined) / self.total_rows if self.total_rows else 0.0

    def reasons(self) -> dict[str, int]:
        counts: dict[str, int] = {}
        for row in self.quarantined:
            counts[row.reason] = counts.get(row.reason, 0) + 1
        return dict(sorted(counts.items(), key=lambda kv: (-kv[1], kv[0])))


@dataclass
class _FormulaIndex:
    """Índice perezoso de celdas que contienen fórmulas.

    `openpyxl` con `data_only=True` devuelve el último valor cacheado por Excel.
    Si un fichero se editó con una herramienta que no recalculó, esas celdas
    llegan vacías. Distinguir "celda vacía" de "fórmula sin valor cacheado" es lo
    que permite reportarlo como incidencia de calidad en vez de inventar un 0
    (docs/00, P3).
    """

    source: bytes
    sheet_name: str
    _index: set[tuple[int, int]] | None = field(default=None, init=False)

    def has_formula(self, row: int, column: int) -> bool:
        if self._index is None:
            self._index = self._build()
        return (row, column) in self._index

    def _build(self) -> set[tuple[int, int]]:
        found: set[tuple[int, int]] = set()
        wb = load_workbook(io.BytesIO(self.source), read_only=True, data_only=False)
        try:
            sheet = wb[self.sheet_name]
            for r, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                for c, value in enumerate(row, start=1):
                    if isinstance(value, str) and value.startswith("="):
                        found.add((r, c))
        finally:
            wb.close()
        return found


def _read_source(source: str | Path | bytes | BinaryIO) -> bytes:
    if isinstance(source, bytes):
        return source
    if isinstance(source, (str, Path)):
        return Path(source).read_bytes()
    return source.read()


def _select_sheet(workbook: Workbook, mapping: SchemaMapping) -> Worksheet:
    wanted = {normalize_key(n) for n in mapping.sheet_names}
    for name in workbook.sheetnames:
        if normalize_key(name) in wanted:
            return workbook[name]
    if len(workbook.sheetnames) == 1:
        return workbook[workbook.sheetnames[0]]
    raise SchemaMappingError(
        f"El contrato '{mapping.name}' espera una hoja llamada {list(mapping.sheet_names)}, "
        f"y el fichero tiene {workbook.sheetnames}. Añade el nombre real a "
        f"mappings/{mapping.name}.yaml en sheet.names."
    )


def _scan_rows(sheet: Worksheet, limit: int) -> list[tuple[int, tuple[Any, ...]]]:
    scanned: list[tuple[int, tuple[Any, ...]]] = []
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if index > limit:
            break
        scanned.append((index, row))
    return scanned


def _find_header(
    scanned: Sequence[tuple[int, tuple[Any, ...]]], mapping: SchemaMapping
) -> tuple[int, dict[int, str], list[str]]:
    """Localiza la fila de cabecera.

    Devuelve (número de fila, columna→campo canónico, cabeceras sin mapear).
    La cabecera rara vez está en la fila 1 en ficheros reales.
    """
    best: tuple[int, int, dict[int, str], list[str]] | None = None

    for row_number, values in scanned:
        assignments: dict[int, str] = {}
        unmapped: list[str] = []
        for column_index, cell in enumerate(values, start=1):
            if is_blank(cell):
                continue
            spec = mapping.column_for_header(cell)
            if spec is None:
                unmapped.append(str(cell).strip())
            elif spec.field not in assignments.values():
                assignments[column_index] = spec.field

        matched_required = sum(
            1 for spec in mapping.required_columns if spec.field in assignments.values()
        )
        if matched_required != len(mapping.required_columns):
            continue
        score = len(assignments)
        if best is None or score > best[1]:
            best = (row_number, score, assignments, unmapped)

    if best is None:
        required = [spec.field for spec in mapping.required_columns]
        raise SchemaMappingError(
            f"No se encontró la fila de cabecera en las primeras {len(scanned)} filas. "
            f"El contrato '{mapping.name}' exige columnas para {required}. "
            f"Revisa los alias en mappings/{mapping.name}.yaml o fija header.row."
        )
    return best[0], best[2], best[3]


def _truncate(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value)
    return text if len(text) <= _MAX_RAW_VALUE else text[:_MAX_RAW_VALUE] + "…"


def _looks_like_subtotal(raw: Mapping[str, Any], mapping: SchemaMapping) -> bool:
    """Fila de subtotal: sin identificadores, pero con números.

    Es el patrón típico de las filas de agregación intercaladas en los planes de
    proyecto reales. Se reporta con su propio motivo para que no se confunda con
    un fichero mal mapeado.
    """
    text_required = [s.field for s in mapping.required_columns if s.type in ("text", "id_list")]
    if not text_required or any(not is_blank(raw.get(f)) for f in text_required):
        return False
    numeric = [s.field for s in mapping.columns if s.type == "number"]
    return any(not is_blank(raw.get(f)) for f in numeric)


def _coerce_row(
    raw: Mapping[str, Any],
    column_index: Mapping[str, int],
    mapping: SchemaMapping,
    row_number: int,
    formulas: _FormulaIndex,
) -> tuple[dict[str, Any] | None, QuarantinedRow | None]:
    data: dict[str, Any] = {}
    for spec in mapping.columns:
        value = raw.get(spec.field)

        if is_blank(value) and spec.required:
            column = column_index.get(spec.field)
            if column is not None and formulas.has_formula(row_number, column):
                return None, QuarantinedRow(
                    row_number=row_number,
                    reason="formula_without_cached_value",
                    field_name=spec.field,
                    raw_value=None,
                    detail=(
                        "La celda contiene una fórmula sin valor cacheado. Abre el fichero "
                        "en Excel y guárdalo para que recalcule."
                    ),
                )
            return None, QuarantinedRow(
                row_number=row_number,
                reason="missing_required",
                field_name=spec.field,
                raw_value=None,
            )

        try:
            data[spec.field] = COERCERS[spec.type](value)
        except CoercionError as exc:
            return None, QuarantinedRow(
                row_number=row_number,
                reason=exc.reason,
                field_name=spec.field,
                raw_value=_truncate(exc.value),
            )

    return {k: v for k, v in data.items() if v is not None}, None


def _iter_data_rows(sheet: Worksheet, header_row: int) -> Iterator[tuple[int, tuple[Any, ...]]]:
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if index > header_row:
            yield index, row


def parse_workbook(
    source: str | Path | bytes | BinaryIO,
    mapping: SchemaMapping,
) -> ParseResult:
    """Parsea un fichero Excel según su contrato de datos."""
    model = _MODELS.get(mapping.model_name)
    if model is None:
        raise SchemaMappingError(
            f"El contrato '{mapping.name}' declara el modelo desconocido "
            f"'{mapping.model_name}'. Modelos disponibles: {sorted(_MODELS)}"
        )

    payload = _read_source(source)
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        sheet = _select_sheet(workbook, mapping)
        sheet_name = sheet.title

        if mapping.header_detect == "fixed":
            assert mapping.header_row is not None
            scanned = _scan_rows(sheet, mapping.header_row)
            candidates = [r for r in scanned if r[0] == mapping.header_row]
            header_row, assignments, unmapped = _find_header(candidates, mapping)
        else:
            header_row, assignments, unmapped = _find_header(
                _scan_rows(sheet, mapping.max_scan_rows), mapping
            )

        field_to_column = {v: k for k, v in assignments.items()}
        formulas = _FormulaIndex(source=payload, sheet_name=sheet_name)

        rows: list[BaseModel] = []
        quarantined: list[QuarantinedRow] = []

        for row_number, values in _iter_data_rows(sheet, header_row):
            if all(is_blank(v) for v in values):
                continue

            raw = {
                field_name: values[index - 1] if index - 1 < len(values) else None
                for index, field_name in assignments.items()
            }

            if _looks_like_subtotal(raw, mapping):
                quarantined.append(
                    QuarantinedRow(
                        row_number=row_number,
                        reason="probable_subtotal_row",
                        field_name=None,
                        raw_value=None,
                        detail="Fila sin identificador pero con importes: parece un subtotal.",
                    )
                )
                continue

            data, problem = _coerce_row(raw, field_to_column, mapping, row_number, formulas)
            if problem is not None:
                quarantined.append(problem)
                continue

            assert data is not None
            try:
                rows.append(model(**data))
            except ValidationError as exc:
                first = exc.errors()[0]
                quarantined.append(
                    QuarantinedRow(
                        row_number=row_number,
                        reason="validation_error",
                        field_name=str(first.get("loc", ("?",))[0]),
                        raw_value=_truncate(first.get("input")),
                        detail=first.get("msg"),
                    )
                )

        mapped = set(assignments.values())
        missing = tuple(sorted(s.field for s in mapping.columns if s.field not in mapped))

        return ParseResult(
            mapping_name=mapping.name,
            sheet_name=sheet_name,
            header_row=header_row,
            rows=tuple(rows),
            quarantined=tuple(quarantined),
            column_map={v: str(k) for k, v in assignments.items()},
            missing_columns=missing,
            unmapped_headers=tuple(unmapped),
        )
    finally:
        workbook.close()
