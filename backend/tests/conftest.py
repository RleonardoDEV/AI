"""Ficheros Excel de prueba, deliberadamente sucios.

Los planes de proyecto reales no son tablas limpias: la cabecera no está en la
fila 1, hay fechas como texto, filas de subtotal y estados que nadie normalizó.
Estas fixtures reproducen ese desorden a propósito (docs/00, P4).
"""

from __future__ import annotations

import datetime as dt
import io
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def _to_bytes(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def messy_project_plan() -> bytes:
    """Plan de proyecto con cabecera en la fila 7 y varias filas problemáticas."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tareas"  # nombre en español: debe resolverse por alias

    # Filas 1-6: logo, título, metadatos. Ruido antes de la cabecera.
    ws["A1"] = "CONSULTORÍA ACME"
    ws["A2"] = "Project Plan - Project X"
    ws["A4"] = "Cliente:"
    ws["B4"] = "Client A"
    ws["A5"] = "Última actualización:"
    ws["B5"] = dt.datetime(2026, 8, 21)

    header = [
        "Task ID", "Tarea", "Owner", "Planned Start", "Fin Planificado",
        "Actual End", "Estado", "Estimated Hours", "Horas Reales",
        "Hito", "Depends On", "Critica", "Comentarios",
    ]
    ws.append([])  # fila 6 vacía
    for column, value in enumerate(header, start=1):
        ws.cell(row=7, column=column, value=value)

    rows = [
        # ok, retrasada y crítica
        ["T-001", "Análisis funcional", "Ana", dt.datetime(2026, 6, 1), dt.datetime(2026, 6, 20),
         dt.datetime(2026, 6, 28), "Completada", 80, 96, "No", "", "Sí", "cerrada tarde"],
        # fecha como texto en formato europeo
        ["T-002", "Diseño técnico", "Beto", "01/07/2026", "25/07/2026",
         None, "En curso", 120, 140, "no", "T-001", "sí", ""],
        # bloqueada
        ["T-003", "Integración con ERP", "Carla", dt.datetime(2026, 7, 1), dt.datetime(2026, 8, 10),
         None, "Blocked", 160, 60, "NO", "T-002", "SI", "espera al cliente"],
        # otra bloqueada, crítica y retrasada
        ["T-004", "Pruebas de integración", "Dani", dt.datetime(2026, 8, 1),
         dt.datetime(2026, 8, 18),
         None, "on hold", 80, 20, "No", "T-003", "Sí", ""],
        # hito abierto y retrasado
        ["T-005", "UAT start", "Ana", dt.datetime(2026, 8, 15), dt.datetime(2026, 8, 18),
         None, "Pendiente", 0, 0, "Sí", "T-004", "Sí", "hito contractual"],
        # ESTADO DESCONOCIDO -> cuarentena, no se convierte a NOT_STARTED
        ["T-006", "Formación", "Beto", dt.datetime(2026, 9, 1), dt.datetime(2026, 9, 5),
         None, "Semi-hecho", 16, 0, "No", "", "No", ""],
        # FECHA IMPOSIBLE DE PARSEAR -> cuarentena
        ["T-007", "Documentación", "Carla", None, "cuando se pueda",
         None, "Pendiente", 24, 0, "No", "", "No", ""],
        # FILA DE SUBTOTAL -> cuarentena con su propio motivo
        [None, None, None, None, None, None, None, 480, 316, None, None, None, "TOTAL"],
        # fila completamente vacía -> se ignora, no cuenta como error
        [None, None, None, None, None, None, None, None, None, None, None, None, None],
        # tarea sin fecha de fin planificada (obligatoria) -> cuarentena
        ["T-008", "Soporte post go-live", "Dani", dt.datetime(2026, 10, 1), None,
         None, "Pendiente", 40, 0, "No", "", "No", ""],
    ]
    for row in rows:
        ws.append(row)

    return _to_bytes(wb)


@pytest.fixture
def plan_with_stale_formula() -> bytes:
    """Fichero cuya fecha de fin es una fórmula sin valor cacheado."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tasks"
    ws.append(["Task ID", "Task", "Planned End", "Status"])
    ws.append(["T-001", "Tarea normal", dt.datetime(2026, 6, 20), "Done"])
    ws.append(["T-002", "Tarea calculada", "=C2+30", "In Progress"])
    return _to_bytes(wb)


@pytest.fixture
def clean_resources() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Recursos"
    ws.append(["Recurso", "Rol", "Dedicacion", "Dias Ausencia", "Recurso Clave"])
    ws.append(["Ana", "Analista", 100, 0, "No"])
    ws.append(["Beto", "Developer", 80, 0, "No"])
    ws.append(["Carla", "Tech Lead", 120, 5, "Sí"])
    ws.append(["Dani", "Developer", 70, 0, "No"])
    return _to_bytes(wb)


@pytest.fixture
def clean_budget() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Presupuesto"
    ws.append(["Concepto", "Importe Planificado", "Importe Real", "Moneda"])
    ws.append(["Desarrollo", 120000, 141000, "EUR"])
    ws.append(["Licencias", 30000, 30000, "EUR"])
    return _to_bytes(wb)


@pytest.fixture
def risks_book() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Riesgos"
    ws.append(["Risk ID", "Descripcion", "Probabilidad", "Impacto", "Estado", "Responsable"])
    ws.append(["R-01", "Dependencia de integración externa", "Alta", "Alto", "Abierto", "Ana"])
    ws.append(["R-02", "Disponibilidad del Tech Lead", "Media", "Alto", "Mitigando", "Ana"])
    ws.append(["R-03", "Cambio de alcance del cliente", "Baja", "Medio", "Cerrado", "Beto"])
    return _to_bytes(wb)
